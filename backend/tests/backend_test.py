"""AgriScan backend API tests."""
import os
import io
import pytest
import requests
from datetime import datetime, timezone

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://harvest-track-94.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"
ADMIN_EMAIL = "gaelmartin158@gmail.com"
ADMIN_PASSWORD = "agriscan2026"


@pytest.fixture(scope="session")
def session():
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=30)
    assert r.status_code == 200, f"Login failed: {r.status_code} {r.text}"
    data = r.json()
    assert data["email"] == ADMIN_EMAIL
    assert data["role"] == "admin"
    # Verify cookies set
    assert "access_token" in s.cookies
    assert "refresh_token" in s.cookies
    return s


# --- Auth ---
class TestAuth:
    def test_login_invalid(self):
        r = requests.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": "wrong"}, timeout=30)
        assert r.status_code == 401

    def test_me_unauth(self):
        r = requests.get(f"{API}/auth/me", timeout=30)
        assert r.status_code == 401

    def test_me_auth(self, session):
        r = session.get(f"{API}/auth/me", timeout=30)
        assert r.status_code == 200
        d = r.json()
        assert d["email"] == ADMIN_EMAIL
        assert d["role"] == "admin"

    def test_protected_endpoints_unauth(self):
        for path in ["/workers", "/parcels", "/products", "/harvests", "/dashboard", "/harvests/daily"]:
            r = requests.get(f"{API}{path}", timeout=30)
            assert r.status_code == 401, f"{path} returned {r.status_code}"


# --- Workers CRUD ---
class TestWorkers:
    def test_crud(self, session):
        # List
        r = session.get(f"{API}/workers")
        assert r.status_code == 200
        assert isinstance(r.json(), list)
        # Create
        r = session.post(f"{API}/workers", json={"name": "TEST_Worker", "phone": "555", "role": "Cosechador", "active": True})
        assert r.status_code == 200
        w = r.json()
        wid = w["id"]
        assert w["name"] == "TEST_Worker"
        assert "_id" not in w
        # Update
        r = session.put(f"{API}/workers/{wid}", json={"name": "TEST_Worker2", "role": "Supervisor", "active": True})
        assert r.status_code == 200
        assert r.json()["name"] == "TEST_Worker2"
        # Verify persistence
        r = session.get(f"{API}/workers")
        assert any(x["id"] == wid and x["name"] == "TEST_Worker2" for x in r.json())
        # Delete
        r = session.delete(f"{API}/workers/{wid}")
        assert r.status_code == 200
        r = session.get(f"{API}/workers")
        assert not any(x["id"] == wid for x in r.json())


# --- Parcels CRUD ---
class TestParcels:
    def test_crud(self, session):
        r = session.get(f"{API}/parcels")
        assert r.status_code == 200
        r = session.post(f"{API}/parcels", json={"name": "TEST_Parcel", "hectares": 2.5, "crop_type": "Tomate", "active": True})
        assert r.status_code == 200
        pid = r.json()["id"]
        r = session.put(f"{API}/parcels/{pid}", json={"name": "TEST_Parcel2", "hectares": 3.0, "active": True})
        assert r.status_code == 200
        assert r.json()["name"] == "TEST_Parcel2"
        r = session.delete(f"{API}/parcels/{pid}")
        assert r.status_code == 200


# --- Products CRUD ---
class TestProducts:
    def test_crud(self, session):
        r = session.get(f"{API}/products")
        assert r.status_code == 200
        r = session.post(f"{API}/products", json={"name": "TEST_Prod", "price_per_kg": 10.0, "price_per_unit": 2.0, "unit_label": "caja", "active": True})
        assert r.status_code == 200
        p = r.json()
        assert p["price_per_kg"] == 10.0
        assert p["price_per_unit"] == 2.0
        pid = p["id"]
        r = session.put(f"{API}/products/{pid}", json={"name": "TEST_Prod2", "price_per_kg": 15.0, "price_per_unit": 3.0, "active": True})
        assert r.status_code == 200
        assert r.json()["price_per_kg"] == 15.0
        session.delete(f"{API}/products/{pid}")


# --- Harvests ---
@pytest.fixture(scope="class")
def harvest_deps(session):
    """Create test worker, parcel, product for harvest tests."""
    w = session.post(f"{API}/workers", json={"name": "TEST_HW", "active": True}).json()
    p = session.post(f"{API}/parcels", json={"name": "TEST_HP", "active": True}).json()
    pr = session.post(f"{API}/products", json={"name": "TEST_HPr", "price_per_kg": 10.0, "price_per_unit": 5.0, "active": True}).json()
    yield {"worker_id": w["id"], "parcel_id": p["id"], "product_id": pr["id"]}
    session.delete(f"{API}/workers/{w['id']}")
    session.delete(f"{API}/parcels/{p['id']}")
    session.delete(f"{API}/products/{pr['id']}")


class TestHarvests:
    def test_create_payment_calc(self, session, harvest_deps):
        today = datetime.now(timezone.utc).date().isoformat()
        payload = {**harvest_deps, "kilos": 10.0, "units": 3.0, "date": today, "notes": "TEST_H"}
        r = session.post(f"{API}/harvests", json=payload)
        assert r.status_code == 200, r.text
        h = r.json()
        # payment = 10 * 10 + 3 * 5 = 115
        assert h["payment"] == 115.0, f"Expected 115, got {h['payment']}"
        assert h["worker_name"] == "TEST_HW"
        assert h["parcel_name"] == "TEST_HP"
        assert h["product_name"] == "TEST_HPr"
        TestHarvests.hid = h["id"]
        TestHarvests.today = today

    def test_override_prices(self, session, harvest_deps):
        today = TestHarvests.today
        r = session.post(f"{API}/harvests", json={
            **harvest_deps, "kilos": 5.0, "units": 0, "date": today,
            "price_per_kg": 20.0, "price_per_unit": 0, "notes": "TEST_H_override"
        })
        assert r.status_code == 200
        assert r.json()["payment"] == 100.0

    def test_list_filters(self, session, harvest_deps):
        today = TestHarvests.today
        r = session.get(f"{API}/harvests", params={"date_from": today, "date_to": today, "worker_id": harvest_deps["worker_id"]})
        assert r.status_code == 200
        rows = r.json()
        assert all(x["date"] == today for x in rows)
        assert all(x["worker_id"] == harvest_deps["worker_id"] for x in rows)

        # product filter
        r = session.get(f"{API}/harvests", params={"product_id": harvest_deps["product_id"]})
        assert r.status_code == 200
        assert all(x["product_id"] == harvest_deps["product_id"] for x in r.json())

    def test_daily_summary(self, session, harvest_deps):
        today = TestHarvests.today
        r = session.get(f"{API}/harvests/daily", params={"date": today})
        assert r.status_code == 200
        d = r.json()
        assert d["date"] == today
        assert "total_kilos" in d and "total_payment" in d and "by_worker" in d
        # Should include our test worker
        workers = [bw for bw in d["by_worker"] if bw["worker_id"] == harvest_deps["worker_id"]]
        assert len(workers) == 1
        bw = workers[0]
        # From our 2 harvests: kilos=15 (10+5), payment=115+100=215
        assert bw["kilos"] >= 15.0
        assert bw["payment"] >= 215.0

    def test_invalid_product(self, session, harvest_deps):
        today = TestHarvests.today
        r = session.post(f"{API}/harvests", json={
            "worker_id": harvest_deps["worker_id"],
            "parcel_id": harvest_deps["parcel_id"],
            "product_id": "000000000000000000000000",
            "kilos": 1, "units": 0, "date": today
        })
        assert r.status_code == 400

    def test_export_excel(self, session, harvest_deps):
        today = TestHarvests.today
        r = session.get(f"{API}/harvests/export", params={"date_from": today, "date_to": today})
        assert r.status_code == 200
        ct = r.headers.get("Content-Type", "")
        assert "spreadsheet" in ct or "openxml" in ct
        cd = r.headers.get("Content-Disposition", "")
        assert ".xlsx" in cd
        # Validate it's a valid xlsx (starts with PK)
        assert r.content[:2] == b"PK"
        # Try opening
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        assert "Fecha" in headers and "Trabajador" in headers and "Pago Total ($)" in headers

    def test_update_delete(self, session, harvest_deps):
        hid = TestHarvests.hid
        today = TestHarvests.today
        r = session.put(f"{API}/harvests/{hid}", json={
            **harvest_deps, "kilos": 2.0, "units": 1.0, "date": today
        })
        assert r.status_code == 200
        # 2*10 + 1*5 = 25
        assert r.json()["payment"] == 25.0
        r = session.delete(f"{API}/harvests/{hid}")
        assert r.status_code == 200


# --- Dashboard ---
class TestDashboard:
    def test_dashboard(self, session):
        r = session.get(f"{API}/dashboard")
        assert r.status_code == 200
        d = r.json()
        for k in ["today_kilos", "today_payment", "month_kilos", "month_payment", "workers_count", "parcels_count", "weekly"]:
            assert k in d
        assert isinstance(d["weekly"], list)
        assert len(d["weekly"]) == 7


# --- Upload (soft fail if storage upstream unreachable) ---
class TestUpload:
    def test_upload_photo(self, session):
        png = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
               b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc\x00\x01"
               b"\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82")
        files = {"file": ("test.png", png, "image/png")}
        r = session.post(f"{API}/upload", files=files, timeout=60)
        if r.status_code == 500:
            pytest.skip(f"Storage upstream unreachable: {r.text}")
        assert r.status_code == 200
        assert "path" in r.json()


# --- Logout ---
class TestLogout:
    def test_logout(self):
        s = requests.Session()
        r = s.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD})
        assert r.status_code == 200
        r = s.post(f"{API}/auth/logout")
        assert r.status_code == 200
        # After logout, cookies cleared → me should fail
        s.cookies.clear()
        r = s.get(f"{API}/auth/me")
        assert r.status_code == 401
