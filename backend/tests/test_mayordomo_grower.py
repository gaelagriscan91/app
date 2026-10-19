# Mayordomo/Grower fields: check-in session, harvest create/update, harvests & receipts Excel exports
import io
import os
import pytest
import requests
from datetime import date
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
EMAIL = os.environ.get("TEST_ADMIN_EMAIL")
PASSWORD = os.environ.get("TEST_ADMIN_PASSWORD")
TODAY = date.today().isoformat()


@pytest.fixture(scope="module")
def s():
    sess = requests.Session()
    r = sess.post(f"{BASE_URL}/api/auth/login", json={"email": EMAIL, "password": PASSWORD})
    if r.status_code != 200:
        pytest.skip("login failed")
    return sess


@pytest.fixture(scope="module")
def ctx(s):
    orig = s.get(f"{BASE_URL}/api/checkin-session").json()
    parcels = s.get(f"{BASE_URL}/api/parcels").json()
    products = s.get(f"{BASE_URL}/api/products").json()
    assert parcels and products
    w = s.post(f"{BASE_URL}/api/workers", json={"name": "TEST_MG Worker"}).json()
    c = {"orig": orig, "parcel": parcels[0], "product": products[0], "worker": w, "harvests": []}
    yield c
    for hid in c["harvests"]:
        s.delete(f"{BASE_URL}/api/harvests/{hid}")
    s.delete(f"{BASE_URL}/api/workers/{w['id']}")
    s.put(f"{BASE_URL}/api/checkin-session", json=orig)


def test_session_persists_mayordomo_grower(s, ctx):
    data = {**ctx["orig"], "mayordomo": "TEST_May", "grower": "TEST_Grow"}
    assert s.put(f"{BASE_URL}/api/checkin-session", json=data).status_code == 200
    got = s.get(f"{BASE_URL}/api/checkin-session").json()
    assert got["mayordomo"] == "TEST_May" and got["grower"] == "TEST_Grow"


def test_create_harvest_stores_fields(s, ctx):
    r = s.post(f"{BASE_URL}/api/harvests", json={
        "worker_id": ctx["worker"]["id"], "parcel_id": ctx["parcel"]["id"], "product_id": ctx["product"]["id"],
        "kilos": 0, "units": 2, "price_per_kg": 0, "price_per_unit": 1.5, "price_mode": "piece", "unit_type": "pieza",
        "date": TODAY, "notes": "TEST_", "mayordomo": "TEST_May", "grower": "TEST_Grow"})
    assert r.status_code in (200, 201), r.text
    h = r.json()
    ctx["harvests"].append(h["id"])
    assert h["mayordomo"] == "TEST_May" and h["grower"] == "TEST_Grow"
    rows = s.get(f"{BASE_URL}/api/harvests", params={"date_from": TODAY, "date_to": TODAY}).json()
    mine = [x for x in rows if x["id"] == h["id"]]
    assert mine and mine[0]["grower"] == "TEST_Grow" and mine[0]["mayordomo"] == "TEST_May"


def test_update_harvest_keeps_grower_when_omitted(s, ctx):
    hid = ctx["harvests"][0]
    r = s.put(f"{BASE_URL}/api/harvests/{hid}", json={
        "worker_id": ctx["worker"]["id"], "parcel_id": ctx["parcel"]["id"], "product_id": ctx["product"]["id"],
        "kilos": 0, "units": 3, "price_per_kg": 0, "price_per_unit": 1.5, "date": TODAY, "mayordomo": "TEST_May"})
    assert r.status_code == 200, r.text
    assert r.json()["grower"] == "TEST_Grow"
    assert r.json()["units"] == 3


def test_receipts_include_grower(s, ctx):
    rows = s.get(f"{BASE_URL}/api/receipts", params={"date_from": TODAY, "date_to": TODAY}).json()
    mine = [r for r in rows if r.get("worker_id") == ctx["worker"]["id"]]
    assert mine
    assert mine[0]["harvests"][0]["grower"] == "TEST_Grow"


def test_harvests_export_columns(s, ctx):
    r = s.get(f"{BASE_URL}/api/harvests/export", params={"date_from": TODAY, "date_to": TODAY})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    hdr = [c.value for c in ws[1]]
    assert "Mayordomo" in hdr and "Grower" in hdr
    gi = hdr.index("Grower")
    vals = [row[gi].value for row in ws.iter_rows(min_row=2) if row[1].value == "TEST_MG Worker"]
    assert "TEST_Grow" in vals


def test_receipts_export_header(s, ctx):
    r = s.get(f"{BASE_URL}/api/receipts/export", params={"date_from": TODAY, "date_to": TODAY, "tz_offset": 0})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    row3 = [str(c.value or "") for c in ws[3]]
    print(row3)
    assert ws["G3"].value and "Grower" in str(ws["G3"].value)
    assert any("Mayordomo" in v for v in row3)
    assert any("TEST_Grow" in v for v in row3)
    assert any("TEST_May" in v for v in row3)
