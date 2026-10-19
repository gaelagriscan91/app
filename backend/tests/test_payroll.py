# Weekly payroll (Lunes..Domingo per grower): JSON aggregation, dedupe, grower filter, xlsx export
import os
import io
import datetime
import pytest
import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

load_dotenv("/app/backend/.env")
BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/") + "/api"
EMAIL = os.environ.get("TEST_ADMIN_EMAIL")
PASSWORD = os.environ.get("TEST_ADMIN_PASSWORD")
TODAY = datetime.date.today()
MONDAY = TODAY - datetime.timedelta(days=TODAY.weekday())
D = [(MONDAY + datetime.timedelta(days=i)).isoformat() for i in range(7)]
G = "TEST_Grower"


@pytest.fixture(scope="module")
def client():
    s = requests.Session()
    r = s.post(f"{BASE_URL}/auth/login", json={"email": EMAIL, "password": PASSWORD})
    if r.status_code != 200:
        pytest.skip("login failed")
    s.rids = []
    yield s
    for rid in s.rids:
        s.delete(f"{BASE_URL}/receipts/history/{rid}")
        s.delete(f"{BASE_URL}/receipts/trash/{rid}")
    for w in s.get(f"{BASE_URL}/workers").json():
        if w["name"].startswith("TEST_PAY"):
            for h in s.get(f"{BASE_URL}/harvests", params={"worker_id": w["id"]}).json():
                if h.get("worker_id") == w["id"]:
                    s.delete(f"{BASE_URL}/harvests/{h['id']}")
            s.delete(f"{BASE_URL}/workers/{w['id']}")


def _harvest(c, wid, parcel, product, date, units, price, grower, mayordomo="TEST_May"):
    r = c.post(f"{BASE_URL}/harvests", json={
        "worker_id": wid, "parcel_id": parcel, "product_id": product, "kilos": 0, "units": units,
        "price_per_kg": 0, "price_per_unit": price, "price_mode": "piece", "unit_type": "pieza", "date": date,
        "photo_path": "", "notes": "", "block": "", "variety": "", "mayordomo": mayordomo, "grower": grower})
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture(scope="module")
def seed(client):
    c = client
    sess = c.get(f"{BASE_URL}/checkin-session").json()
    parcel = sess.get("parcelId") or c.get(f"{BASE_URL}/parcels").json()[0]["id"]
    product = sess.get("productId") or c.get(f"{BASE_URL}/products").json()[0]["id"]
    a = c.post(f"{BASE_URL}/workers", json={"name": "TEST_PAY Ana"}).json()
    b = c.post(f"{BASE_URL}/workers", json={"name": "TEST_PAY Beto"}).json()
    _harvest(c, a["id"], parcel, product, D[0], 10, 2, G)       # Mon 20
    _harvest(c, a["id"], parcel, product, D[2], 5, 2, G)        # Wed 10
    _harvest(c, b["id"], parcel, product, D[0], 4, 2, G)        # Mon 8
    _harvest(c, b["id"], parcel, product, D[6], 3, 3, "TEST_OtherGrower")  # Sun 9 other grower
    # two overlapping saved receipts that include the same harvests -> must not double count
    for body in ({"date_from": D[0], "date_to": D[6], "worker_id": a["id"]},
                 {"date_from": D[0], "date_to": D[2], "worker_id": a["id"]}):
        r = c.post(f"{BASE_URL}/receipts/history", json=body)
        assert r.status_code == 200, r.text
        c.rids.append(r.json()["id"])
    return {"a": a, "b": b}


def _get(c, **params):
    r = c.get(f"{BASE_URL}/payroll", params=params)
    assert r.status_code == 200, r.text
    return r.json()


def _row(data, name):
    return next((r for r in data["rows"] if r["worker_name"] == name), None)


def test_week_structure(client, seed):
    data = _get(client, week=D[3], grower=G)  # Thursday -> normalized to Monday
    assert data["week_start"] == D[0] and data["week_end"] == D[6]
    assert [d["name"] for d in data["days"]] == ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
    assert [d["date"] for d in data["days"]] == D
    assert data["days"][0]["price"] == 2.0
    assert G in _get(client, week=D[0])["growers"]
    assert "TEST_May" in data["mayordomo"]


def test_grower_filter_and_no_double_count(client, seed):
    data = _get(client, week=D[0], grower=G)
    ana = _row(data, "TEST_PAY Ana")
    beto = _row(data, "TEST_PAY Beto")
    assert ana and beto
    assert ana["days"][0] == {"units": 10.0, "payment": 20.0}
    assert ana["days"][2] == {"units": 5.0, "payment": 10.0}
    assert ana["total"] == 30.0  # not 60/90 even though in 2 receipts + live
    assert beto["total"] == 8.0 and beto["days"][6]["units"] == 0
    assert all(r["worker_name"].startswith("TEST_PAY") for r in data["rows"])
    assert data["total"] == 38.0


def test_grower_case_insensitive(client, seed):
    data = _get(client, week=D[0], grower="test_grower")
    assert data["total"] == 38.0


def test_all_growers(client, seed):
    data = _get(client, week=D[0])
    beto = _row(data, "TEST_PAY Beto")
    assert beto["total"] == 17.0 and beto["days"][6] == {"units": 3.0, "payment": 9.0}


def test_other_week_empty_for_grower(client, seed):
    data = _get(client, week=(MONDAY - datetime.timedelta(days=70)).isoformat(), grower=G)
    assert data["rows"] == [] and data["total"] == 0


def test_invalid_week_and_auth(client):
    assert client.get(f"{BASE_URL}/payroll", params={"week": "bad"}).status_code == 400
    assert requests.get(f"{BASE_URL}/payroll", params={"week": D[0]}).status_code == 401
    assert requests.get(f"{BASE_URL}/payroll/export", params={"week": D[0]}).status_code == 401


def test_export_xlsx(client, seed):
    r = client.get(f"{BASE_URL}/payroll/export", params={"week": D[0], "grower": G})
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert f"payroll_{D[0]}.xlsx" in r.headers.get("content-disposition", "")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A1"].value == "Grower" and ws["B1"].value == G
    assert ws["F1"].value == "Mayordomo" and "TEST_May" in (ws["H1"].value or "")
    assert ws.cell(3, 2).value == "Lunes" and ws.cell(3, 14).value == "Domingo"
    assert ws.cell(4, 2).value == MONDAY.strftime("%m-%d-%y")
    assert ws.cell(5, 2).value == "Pieza" and ws.cell(5, 3).value == "Precio $2.00"
    assert ws.cell(3, 16).value == "$"
    names = {ws.cell(i, 1).value: i for i in range(6, ws.max_row + 1)}
    ri = names["TEST_PAY Ana"]
    assert ws.cell(ri, 2).value == 10 and ws.cell(ri, 3).value == 20 and ws.cell(ri, 16).value == 30
